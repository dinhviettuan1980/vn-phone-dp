from __future__ import annotations

from typing import Iterator

from openpyxl import load_workbook

from dataset_ingestion.base_ingestor import BaseDatasetIngestor


class XlsxIngestor(BaseDatasetIngestor):
    def _sheet(self):
        wb = load_workbook(self.file_path, read_only=True, data_only=True)
        return wb[wb.sheetnames[0]]

    def validate(self) -> None:
        sheet = self._sheet()
        first_row = next(sheet.iter_rows(max_row=1, values_only=True), None)
        if not first_row:
            raise ValueError("XLSX sheet has no header row")

    def headers(self) -> list[str]:
        sheet = self._sheet()
        first_row = next(sheet.iter_rows(max_row=1, values_only=True))
        return [str(h) if h is not None else "" for h in first_row]

    def read_rows(self) -> Iterator[dict]:
        sheet = self._sheet()
        rows = sheet.iter_rows(values_only=True)
        headers = [str(h) if h is not None else "" for h in next(rows)]
        for row in rows:
            yield {headers[i]: row[i] for i in range(len(headers)) if i < len(row)}
